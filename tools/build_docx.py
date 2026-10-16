"""
把方案书 Markdown 转成排版可用的 .docx。

为什么要这个脚本：pandoc 转出来的 docx 对中文只给西文字体，
在 Word 里中文字体会回退成宋体或不统一。这里统一把 eastAsia 设为微软雅黑，
并给标题加段前间距，避免正文与标题挤在一起。

依赖：pandoc（命令行）、python-docx
用法：python tools/build_docx.py
"""

import os
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

# 竞赛方案属于内部材料，不放进开源仓库，放在仓库外的 01-竞赛材料/。
# 可用环境变量 QINGYU_DOC_DIR 覆盖路径。
DOC_DIR = Path(os.environ.get('QINGYU_DOC_DIR') or (ROOT.parent.parent / '01-竞赛材料'))
SRC = DOC_DIR / '轻语TTS-软件创新大赛项目方案.md'
DST = DOC_DIR / '轻语TTS-软件创新大赛项目方案.docx'
TMP = Path(__file__).resolve().parent / '_tmp.docx'

CN_FONT = '微软雅黑'


def md_to_docx() -> None:
    r = subprocess.run(
        ['pandoc', str(SRC), '-o', str(TMP), '--from', 'markdown+raw_tex'],
        capture_output=True, text=True, encoding='utf-8',
    )
    if r.returncode != 0:
        raise SystemExit('pandoc 失败：\n' + (r.stderr or r.stdout))
    print('[1/3] pandoc 已生成中间文件 %.1f KB' % (TMP.stat().st_size / 1024))


def restyle() -> None:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document(str(TMP))

    normal = doc.styles['Normal']
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)

    # 标题：统一中文字体，并把段前拉开
    for name, before in (('Heading 1', 16), ('Heading 2', 12), ('Heading 3', 8)):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = CN_FONT
        st.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
        st.paragraph_format.space_before = Pt(before)

    # 正文中的西文字体也被 Normal 带过去了，但部分 run 自带字体要单独刷
    for p in doc.paragraphs:
        for run in p.runs:
            run.font.name = CN_FONT
            if run._element.rPr is not None:
                run._element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = CN_FONT
                        if run._element.rPr is not None:
                            run._element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)

    doc.save(str(DST))

    paras = len(doc.paragraphs)
    tables = len(doc.tables)
    print('[2/3] 已统一中文字体为 %s' % CN_FONT)
    print('[3/3] 完成 -> %s  （%d 段 / %d 表 / %.1f KB）'
          % (DST.name, paras, tables, DST.stat().st_size / 1024))


def main() -> None:
    if not SRC.exists():
        raise SystemExit('找不到源文件：' + str(SRC))
    md_to_docx()
    try:
        restyle()
    finally:
        if TMP.exists():
            TMP.unlink()
    print('\n完成。')


if __name__ == '__main__':
    sys.exit(main())
